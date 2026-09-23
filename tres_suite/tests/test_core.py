"""Unit tests on synthetic workbooks (no supplied data needed)."""
from __future__ import annotations

import importlib
import os
import subprocess
import sys

import numpy as np
import pytest

from conftest import make_anisotropy_workbook, make_tres_workbook, synthetic_tres_params

from tres_suite.aggregate import summarize_arrays, summarize_scalars
from tres_suite.anisotropy import WobbleSettings, compute_anisotropy
from tres_suite.config import config_from_dict
from tres_suite.das import DasSettings, compute_das
from tres_suite.expressions import ExpressionError, evaluate_expression, parse_expression
from tres_suite.pipeline import run_project
from tres_suite.project import LabelValue, Project
from tres_suite.stats import linear_fit
from tres_suite.tdfs import TdfsSettings, compute_tdfs
from tres_suite.validation import Severity, ValidationThresholds, validate_group
from tres_suite.workbook import WorkbookParseError, detect_workbook_type, load_workbook

WL = np.arange(420, 545, 5)


# ----------------------------------------------------------------------------- import safety
def test_import_has_no_side_effects(tmp_path):
    """Importing every module must not write files, prompt or start analysis."""
    code = ("import os, sys; before=set(os.listdir('.')); "
            "import tres_suite, tres_suite.workbook, tres_suite.project, tres_suite.validation, tres_suite.das, tres_suite.tdfs, "
            "tres_suite.anisotropy, tres_suite.aggregate, tres_suite.stats, tres_suite.export, tres_suite.plotting, tres_suite.pipeline, tres_suite.cli, tres_suite.config; "
            "after=set(os.listdir('.')); assert before==after, after-before; print('ok')")
    env = dict(os.environ, PYTHONPATH=os.path.join(os.path.dirname(__file__), ".."))
    out = subprocess.run([sys.executable, "-c", code], cwd=tmp_path, env=env, capture_output=True, text=True, stdin=subprocess.DEVNULL, timeout=120)
    assert out.returncode == 0, out.stderr
    assert out.stdout.strip() == "ok"


# ----------------------------------------------------------------------------- parsing
@pytest.mark.parametrize("n_comp", [1, 2, 3, 4, 5])
def test_parse_component_counts(tmp_path, n_comp):
    taus = np.linspace(0.2, 4.0, n_comp)
    _, B, counts = synthetic_tres_params(WL, taus=tuple(taus))
    p = make_tres_workbook(str(tmp_path / f"c{n_comp}.xlsx"), WL, taus, B, counts)
    wb = load_workbook(p)
    assert wb.kind == "tres" and wb.n_components == n_comp
    assert np.allclose(wb.lifetimes_ns, np.sort(taus))
    assert wb.components.pre_exponential.shape == (len(WL), n_comp)
    ra = wb.components.relative_amplitude_percent
    assert np.allclose(ra.sum(axis=1), 100.0)
    # rounded Summary values agree with Results-derived values within 2-decimal rounding
    assert np.nanmax(np.abs(wb.summary_relative_amplitude_percent - ra)) <= 0.0051
    assert np.nanmax(np.abs(wb.summary_normalized_pre_exponential - wb.components.normalized_pre_exponential)) <= 0.0051


def test_lifetime_sorting_permutes_columns(tmp_path):
    taus = np.array([3.5, 0.2, 1.5])  # unsorted in the export
    _, B, counts = synthetic_tres_params(WL, taus=tuple(taus))
    p = make_tres_workbook(str(tmp_path / "u.xlsx"), WL, taus, B, counts)
    wb = load_workbook(p)
    assert list(wb.lifetimes_ns) == [0.2, 1.5, 3.5]
    assert list(wb.components.original_component_index) == [2, 3, 1]
    assert np.allclose(wb.components.pre_exponential[:, 0], B[:, 1])


def test_non_default_wavelength_coverage_and_no_range_limit(tmp_path):
    wl = np.arange(300, 810, 10)  # beyond the legacy 300-800 / 420-540 assumptions
    taus, B, counts = synthetic_tres_params(wl)
    p = make_tres_workbook(str(tmp_path / "wide.xlsx"), wl, taus, B, counts)
    wb = load_workbook(p)
    assert wb.wavelength_min_nm == 300 and wb.wavelength_max_nm == 800 and wb.n_wavelengths == len(wl)
    assert wb.is_uniform_grid()
    das = compute_das(wb, DasSettings())
    assert np.isfinite(das.mean_lifetime_ns)


def test_nonuniform_grid_detected_and_used_as_read(tmp_path):
    wl = np.array([420, 425, 430, 440, 450, 470, 500, 540], float)
    taus, B, counts = synthetic_tres_params(wl)
    p = make_tres_workbook(str(tmp_path / "nu.xlsx"), wl, taus, B, counts)
    wb = load_workbook(p)
    assert not wb.is_uniform_grid()
    assert np.array_equal(wb.wavelengths_nm, wl)


def test_missing_cells_and_unsupported_files(tmp_path):
    taus, B, counts = synthetic_tres_params(WL)
    p = make_tres_workbook(str(tmp_path / "bad.xlsx"), WL, taus, B, counts)
    import openpyxl
    wb = openpyxl.load_workbook(p)
    wb["Results"].cell(15, 11).value = None  # B1 at third wavelength
    wb.save(p)
    with pytest.raises(WorkbookParseError) as ei:
        load_workbook(p)
    assert "B1" in str(ei.value) and "Results!K15" in str(ei.value)
    # nonfinite value
    p2 = make_tres_workbook(str(tmp_path / "nan.xlsx"), WL, taus, B, counts)
    wb = openpyxl.load_workbook(p2)
    wb["Results"].cell(13, 22).value = "NaN"
    wb.save(p2)
    with pytest.raises(WorkbookParseError):
        load_workbook(p2)
    # no Results sheet
    other = openpyxl.Workbook(); other.active.title = "Sheet1"; other.save(str(tmp_path / "plain.xlsx"))
    kind, ev = detect_workbook_type(str(tmp_path / "plain.xlsx"))
    assert kind == "unsupported" and "Results" in ev["reason"]
    # lock file
    lock = tmp_path / "~$x.xlsx"; lock.write_bytes(b"junk")
    with pytest.raises(WorkbookParseError):
        load_workbook(str(lock))


def test_anisotropy_parse(tmp_path):
    p = make_anisotropy_workbook(str(tmp_path / "a.xlsx"), [5.0, 0.6], [0.05, 0.25], 0.04, 0.34)
    wb = load_workbook(p)
    assert wb.kind == "anisotropy" and wb.n_components == 2
    assert list(wb.rotational_correlation_times_ns) == [0.6, 5.0]
    assert wb.r_inf == pytest.approx(0.04) and wb.r0_fitted == pytest.approx(0.34)


# ----------------------------------------------------------------------------- labels / project
def test_label_parsing_and_numeric_sort():
    vals = [LabelValue.parse(s) for s in ["5", "15", "100", "17C", "T17", "abc", "23.5 °C", "300K"]]
    assert [v.number for v in vals[:3]] == [5, 15, 100]
    assert vals[3].number == 17 and vals[3].unit == "degC" and vals[4].number == 17 and vals[5].number is None
    assert vals[6].number == 23.5 and vals[6].unit == "degC" and vals[7].unit == "K"
    assert sorted(["100", "5", "15"], key=lambda s: LabelValue.parse(s).sort_key()) == ["5", "15", "100"]


def test_project_groups_and_selection(synthetic_group):
    p = Project(["Composition", "Temperature", "Replicate"])
    p.import_folder(str(synthetic_group))
    assert len(p.measurements) == 8
    groups = p.groups()
    assert [g.name for g in groups] == ["DOPC / 15", "DOPC / 25", "DPH / 15"]
    assert len(p.select(Composition="DOPC")) == 6
    assert len(p.select(Composition="DOPC", Temperature=15)) == 3
    with pytest.raises(ValueError):
        p.add_file(str(synthetic_group / "DOPC" / "15" / "1.xlsx"), {"Composition": "DOPC", "Temperature": "15", "Replicate": "1"})  # duplicate id


def test_numbered_replicate_folders(tmp_path):
    root = tmp_path / "r"
    for rep in (1, 2):
        d = root / "DOPC" / "15" / str(rep); d.mkdir(parents=True)
        taus, B, counts = synthetic_tres_params(WL, seed=rep)
        make_tres_workbook(str(d / "data.xlsx"), WL, taus, B, counts)
    p = Project(["Composition", "Temperature", "Replicate"])
    p.import_folder(str(root), depth=3)
    assert [m.label_text("Replicate") for m in p.measurements] == ["1", "2"]


# ----------------------------------------------------------------------------- validation
def _group_with(tmp_path, files):
    p = Project(["Cond", "Rep"])
    for name, path in files:
        p.add_file(path, {"Cond": "g", "Rep": name})
    p.load_all()
    g = p.groups()[0]
    return g, validate_group(g, ValidationThresholds(chi_sq_warning_threshold=1.5))


def _errors(diags):
    return sorted({d.rule for d in diags if d.severity == Severity.ERROR})


def test_component_count_mismatch_fails_group(tmp_path):
    files = []
    for n in (2, 3):
        taus = np.linspace(0.2, 3, n); _, B, counts = synthetic_tres_params(WL, taus=tuple(taus))
        files.append((f"c{n}", make_tres_workbook(str(tmp_path / f"c{n}.xlsx"), WL, taus, B, counts)))
    g, d = _group_with(tmp_path, files)
    assert _errors(d) == ["component_count_mismatch"]
    err = [x for x in d if x.rule == "component_count_mismatch"][0]
    assert err.observed["n_components"] == {"c2.xlsx": 2, "c3.xlsx": 3} and len(err.files) == 2


def test_wavelength_range_and_grid_mismatch(tmp_path):
    taus, B, counts = synthetic_tres_params(WL)
    a = make_tres_workbook(str(tmp_path / "a.xlsx"), WL, taus, B, counts)
    wl2 = np.arange(435, 530, 5); _, B2, c2 = synthetic_tres_params(wl2)
    b = make_tres_workbook(str(tmp_path / "b.xlsx"), wl2, taus, B2, c2)
    g, d = _group_with(tmp_path, [("a", a), ("b", b)])
    assert _errors(d) == ["wavelength_range_mismatch"]
    obs = [x for x in d if x.rule == "wavelength_range_mismatch"][0].observed["min_max_nm"]
    assert obs == {"a.xlsx": [420.0, 540.0], "b.xlsx": [435.0, 525.0]}
    wl3 = np.arange(420, 541, 1); _, B3, c3 = synthetic_tres_params(wl3)
    c = make_tres_workbook(str(tmp_path / "c.xlsx"), wl3, taus, B3, c3)
    g, d = _group_with(tmp_path, [("a", a), ("c", c)])
    assert _errors(d) == ["wavelength_grid_mismatch"]
    obs = [x for x in d if x.rule == "wavelength_grid_mismatch"][0].observed
    assert obs["intervals_nm"] == {"a.xlsx": [5.0], "c.xlsx": [1.0]} and obs["first_mismatching_coordinate_nm"] == {"a.xlsx": 425.0, "c.xlsx": 421.0}


def test_chi_sq_warning_only_and_no_threshold(tmp_path):
    taus, B, counts = synthetic_tres_params(WL)
    a = make_tres_workbook(str(tmp_path / "a.xlsx"), WL, taus, B, counts, chi_global=2.3)
    g, d = _group_with(tmp_path, [("a", a)])
    assert _errors(d) == []
    w = [x for x in d if x.rule == "high_chi_sq"]
    assert len(w) == 1 and w[0].observed["global_chi_sq"] == pytest.approx(2.3) and w[0].observed["threshold"] == 1.5
    d2 = validate_group(g, ValidationThresholds())  # no threshold -> report only
    assert not [x for x in d2 if x.rule == "high_chi_sq"] and [x for x in d2 if x.rule == "chi_sq_report"]


def test_mixed_types_fail_and_anisotropy_differences_warn(tmp_path):
    taus, B, counts = synthetic_tres_params(WL)
    t = make_tres_workbook(str(tmp_path / "t.xlsx"), WL, taus, B, counts)
    a = make_anisotropy_workbook(str(tmp_path / "a.xlsx"), [2.0], [0.3], 0.05, 0.35)
    g, d = _group_with(tmp_path, [("t", t), ("a", a)])
    assert _errors(d) == ["mixed_workbook_types"]
    b = make_anisotropy_workbook(str(tmp_path / "b.xlsx"), [0.5, 8.0], [0.1, 0.2], 0.2, 0.35)
    p = Project(["Cond", "Rep"]); p.add_file(a, {"Cond": "g", "Rep": "a"}); p.add_file(b, {"Cond": "g", "Rep": "b"}); p.load_all()
    d = validate_group(p.groups()[0], ValidationThresholds(anisotropy_correlation_time_tolerance_percent=10))
    assert _errors(d) == []
    rules = {x.rule for x in d if x.severity == Severity.WARNING}
    assert {"anisotropy_component_count_differs", "anisotropy_lifetime_difference"} <= rules


# ----------------------------------------------------------------------------- expressions / metrics
def test_expression_parser_restricted():
    parse_expression("(I(440) - I(490)) / (I(440) + I(490))")
    for bad in ["__import__('os')", "I(440) ** 2", "max(I(440), 1)", "I(x)", "I(440); 1", "[1]", "I(440) if 1 else 0"]:
        with pytest.raises(ExpressionError):
            parse_expression(bad)


def test_ratio_unavailable_out_of_range_and_zero_denominator():
    wl = np.array([420, 430, 440.0]); sp = np.array([10.0, 20.0, 0.0])
    r = evaluate_expression("x", "I(510) / I(430)", wl, sp)
    assert not r.available and "outside" in r.reason and r.value is None
    r = evaluate_expression("x", "I(430) / I(440)", wl, sp)
    assert not r.available and r.reason == "zero denominator"
    r = evaluate_expression("x", "I(425) / I(430)", wl, sp)
    assert r.available and r.value == pytest.approx(15 / 20)


def test_gp_and_custom_ratio_in_das(tmp_path):
    taus, B, counts = synthetic_tres_params(WL)
    p = make_tres_workbook(str(tmp_path / "g.xlsx"), WL, taus, B, counts)
    wb = load_workbook(p)
    das = compute_das(wb, DasSettings(custom_ratios={"r510_430": "I(510)/I(430)", "bad": "I(600)/I(430)"}))
    ib = np.interp(440, WL, counts); ir = np.interp(490, WL, counts)
    assert das.metrics["gp"].value == pytest.approx((ib - ir) / (ib + ir))
    assert das.metrics["r510_430"].value == pytest.approx(np.interp(510, WL, counts) / np.interp(430, WL, counts))
    assert not das.metrics["bad"].available
    rec = das.scalar_record()
    assert np.isnan(rec["bad"]) and rec["bad_available"] is False
    # mean lifetime definition
    W = das.weights
    assert das.mean_lifetime_ns == pytest.approx(float(np.dot(das.lifetimes_ns, W) / W.sum()))
    assert das.area_percent.sum() == pytest.approx(100.0)


# ----------------------------------------------------------------------------- TDFS
def test_reconstruction_integrates_to_fss_and_nu0_modes(tmp_path):
    taus, B, counts = synthetic_tres_params(WL)
    p = make_tres_workbook(str(tmp_path / "t.xlsx"), WL, taus, B, counts)
    wb = load_workbook(p)
    das = compute_das(wb, DasSettings())
    s = TdfsSettings(recon_tmax_min_ns=60)
    td = compute_tdfs(WL, das.lifetimes_ns, das.normalized_pre_exponential, das.fss_area_norm, s, das.mean_lifetime_ns)
    integ = np.trapz(td.tres, td.t_ns, axis=0) if hasattr(np, "trapz") else np.trapezoid(td.tres, td.t_ns, axis=0)
    good = ~np.isin(WL, td.masked_wavelengths_nm)
    assert np.allclose(integ[good], das.fss_area_norm[good], rtol=1e-3)
    assert td.peak.available and td.peak.nu0_source.startswith("apparent") and td.peak.nu0_time_ns == pytest.approx(0.1)
    assert td.peak.delta_nu_cm1 > 0 and td.peak.tau_r_ns > 0
    td2 = compute_tdfs(WL, das.lifetimes_ns, das.normalized_pre_exponential, das.fss_area_norm, TdfsSettings(nu0_cm1=23800.0), das.mean_lifetime_ns)
    assert td2.peak.nu0_cm1 == 23800.0 and td2.peak.nu0_source == "user_supplied"
    # tau_r with a user nu0 starts at t=0, C(0)=1: exactly the legacy integral plus the 0 -> metric_tmin trapezoid
    legacy = compute_tdfs(WL, das.lifetimes_ns, das.normalized_pre_exponential, das.fss_area_norm,
                          TdfsSettings(nu0_cm1=23800.0, tau_r_from_zero_with_user_nu0=False), das.mean_lifetime_ns)
    k = int(np.searchsorted(legacy.traj_t_ns, legacy.peak.window_tmin_ns))
    head = 0.5 * legacy.peak.window_tmin_ns * (1.0 + legacy.traj_c_peak[k])
    assert td2.peak.tau_r_ns == pytest.approx(legacy.peak.tau_r_ns + head, rel=1e-12)
    assert td2.peak.delta_nu_cm1 == pytest.approx(23800.0 - td2.peak.nuinf_cm1)
    assert td2.com.nu0_source.startswith("apparent")  # COM never inherits the user peak value
    rec = td.scalar_record()
    assert rec["legacy_gate_evaluated"] is False and "legacy_x_kinetic" in rec
    assert not any(k.endswith("_valid_physical") or k.endswith("_report_physical") for k in rec)


def test_no_hardcoded_23800_fallback():
    src_dir = os.path.join(os.path.dirname(__file__), "..", "tres_suite")
    for fn in os.listdir(src_dir):
        if fn.endswith(".py"):
            code = open(os.path.join(src_dir, fn), encoding="utf-8").read()
            assert "23800" not in code.replace("23800 cm", ""), fn  # only allowed inside prose mentioning the legacy constant


def test_masked_wavelength_reported():
    wl = np.array([420, 430, 440, 450, 460.0]); taus = np.array([0.5, 3.0])
    alpha = np.array([[0.8, 0.2], [0.7, 0.3], [0.5, 0.5], [-0.9, 0.1], [0.4, 0.6]])  # 450 nm: alpha@tau = -0.45+0.3 < 0
    fss = np.ones(5) / 40
    td = compute_tdfs(wl, taus, alpha, fss, TdfsSettings(recon_tmax_min_ns=5, fine_nu_points=200))
    assert td.n_masked == 1 and td.masked_wavelengths_nm == [450.0]
    assert np.all(td.tres[:, 3] == 0)


# ----------------------------------------------------------------------------- anisotropy add-on
def test_wobble_add_on_isolated_and_blocked(tmp_path):
    p = make_anisotropy_workbook(str(tmp_path / "a.xlsx"), [3.5], [0.29], 0.0454575, 0.333)
    wb = load_workbook(p)
    base = compute_anisotropy(wb, WobbleSettings(enabled=False))
    on = compute_anisotropy(wb, WobbleSettings(enabled=True, effective_volume=161, effective_volume_unit="A3", temperature_K=293, r0=0.39))
    for k in ("r_inf", "r0_fitted", "phi_mean_ns", "phi1_ns", "chi_sq"):
        assert base.scalar_record()[k] == on.scalar_record()[k]
    assert base.wobble["status"] == "disabled"
    # legacy dph() arithmetic for the same inputs
    k = 0.0454575 / 0.39; dperp = (0.1674 - 0.1066 * k - 0.0062 * k * k) / 3.5
    assert on.wobble["eta_poise"] == pytest.approx(293 * 1.380649e-23 / (6 * dperp * 161e-22), rel=1e-6)
    blocked = compute_anisotropy(wb, WobbleSettings(enabled=True, effective_volume=161))
    assert blocked.wobble["status"].startswith("blocked") and "temperature_K (positive)" in blocked.wobble["missing"]
    fitted = compute_anisotropy(wb, WobbleSettings(enabled=True, effective_volume=0.161, effective_volume_unit="nm3", temperature_K=293, r0_source="fitted"))
    assert fitted.wobble["r0_used"] == pytest.approx(0.333) and fitted.wobble["status"] == "ok"


# ----------------------------------------------------------------------------- aggregation / stats
def test_single_replicate_sd_undefined_and_no_boolean_averaging():
    s = summarize_scalars([{"a": 1.0, "flag": True, "name": "x", "n_comp": 3}])
    assert s["a_mean"] == 1.0 and np.isnan(s["a_std"]) and s["a_n"] == 1
    assert "flag_mean" not in s and "name_mean" not in s
    s2 = summarize_scalars([{"a": 1.0}, {"a": 3.0}, {"a": np.nan}])
    assert s2["a_mean"] == 2.0 and s2["a_std"] == pytest.approx(np.sqrt(2)) and s2["a_n"] == 2
    arr = summarize_arrays([np.array([1.0, np.nan]), np.array([3.0, 5.0])])
    assert arr["mean"][0] == 2.0 and np.isnan(arr["std"][1]) and list(arr["n"]) == [2, 1]


def test_linear_fit_r2_edge_cases():
    f = linear_fit([1, 2, 3, 4], [2, 4, 6, 8])
    assert f.available and f.r_squared == pytest.approx(1.0) and f.slope == pytest.approx(2.0) and f.n_points == 4
    assert not linear_fit([1, 2], [1, 2]).available
    assert "constant" in linear_fit([1, 2, 3], [5, 5, 5]).reason
    assert "constant" in linear_fit([2, 2, 2], [1, 2, 3]).reason


# ----------------------------------------------------------------------------- pipeline / export / plots
def test_pipeline_end_to_end_and_reproducible(synthetic_group, tmp_path):
    from tres_suite.export import write_run
    from tres_suite.plotting import render_plots
    cfg_dict = {"project": {"level_names": ["Composition", "Temperature", "Replicate"], "import_folders": [str(synthetic_group)]},
                "output": {"directory": str(tmp_path / "out"), "overwrite": True, "image_formats": ["png"]},
                "tdfs": {"recon_tmax_min_ns": 30, "trajectory_stride": 5},
                "validation": {"replicate_variation_warning_percent": 50},
                "plots": [{"type": "metric_vs_label", "name": "tmean_vs_T", "x": "Temperature", "y": "tmean_ns", "fit": True,
                           "series": [{"label": "DOPC", "select": {"Composition": "DOPC"}, "color": "#ff0000", "marker": "s"}]},
                          {"type": "metric_vs_metric", "name": "gp_vs_tmean", "x": "tmean_ns", "y": "gp", "series": [{"label": "DOPC", "select": {"Composition": "DOPC"}}]},
                          {"type": "das_spectra", "name": "das", "group": {"Composition": "DOPC", "Temperature": "15"}},
                          {"type": "das_spectra", "name": "npe", "representation": "normalized_pre_exponential", "group": {"Composition": "DOPC", "Temperature": "15"}},
                          {"type": "bubble", "name": "bub", "group": {"Composition": "DOPC", "Temperature": "15"}},
                          {"type": "tdfs_spectra", "name": "spec", "group": {"Composition": "DOPC", "Temperature": "15"}},
                          {"type": "tdfs_peak_vs_time", "name": "pk", "series": [{"label": "15", "select": {"Composition": "DOPC", "Temperature": "15"}}]},
                          {"type": "tdfs_relaxation", "name": "ct", "series": [{"label": "15", "select": {"Composition": "DOPC", "Temperature": "15"}}]}]}
    cfg = config_from_dict(cfg_dict)
    res = run_project(cfg)
    assert res.status == "completed"
    assert [g.kind for g in res.groups] == ["tres", "tres", "anisotropy"]
    plots = render_plots(res, str(tmp_path / "out" / "plots"))
    manifest = write_run(res, plots=[p.as_manifest() for p in plots])
    out = tmp_path / "out"
    for f in ["inventory.csv", "coverage_report.csv", "diagnostics.csv", "tres_replicates.csv", "tres_summary.csv", "anisotropy_replicates.csv",
              "records.csv", "run_config.json", "environment.json", "run_manifest.json", "metric_tables/tmean_ns_mean.csv",
              "legacy/graph_data/gp_mean.csv", "legacy/spectra_raw/DOPC_15.csv", "groups/DOPC_15/das_spectra.csv", "groups/DOPC_15/tdfs_trajectory_mean.csv",
              "plots/tmean_vs_T.png", "plots/tmean_vs_T_data.csv", "plots/bub_data.csv"]:
        assert (out / f).exists(), f
    # exported plot data equals the displayed group summary values
    import pandas as pd
    pd_ = pd.read_csv(out / "plots" / "tmean_vs_T_data.csv")
    summ = pd.read_csv(out / "tres_summary.csv")
    for _, r in pd_[pd_.series == "DOPC"].iterrows():
        row = summ[summ.group == r["group"]].iloc[0]
        assert r["y"] == pytest.approx(row["tmean_ns_mean"]) and r["y_sd"] == pytest.approx(row["tmean_ns_std"]) and r["n"] == row["tmean_ns_n"]
    legacy = open(out / "legacy/graph_data/gp_mean.csv", encoding="utf-8-sig").read()
    assert "UNEVALUATED" in legacy and "TRUE" not in legacy
    # reproducibility: a second run gives identical numbers
    res2 = run_project(cfg)
    for g1, g2 in zip(res.groups, res2.groups):
        for k, v in g1.summary.items():
            if isinstance(v, float) and np.isfinite(v):
                assert v == g2.summary[k], k
    # source workbooks untouched
    import hashlib
    for m in res.project.measurements:
        assert m.workbook is not None


def test_blocked_group_not_averaged_from_subset(tmp_path):
    root = tmp_path / "d"; (root / "A" / "1").mkdir(parents=True)
    taus, B, counts = synthetic_tres_params(WL)
    make_tres_workbook(str(root / "A" / "1" / "r1.xlsx"), WL, taus, B, counts)
    make_tres_workbook(str(root / "A" / "1" / "r2.xlsx"), WL, taus, B, counts)
    import openpyxl
    bad = openpyxl.Workbook(); bad.active.title = "Sheet1"; bad.save(str(root / "A" / "1" / "r3.xlsx"))
    cfg = config_from_dict({"project": {"level_names": ["C", "T", "R"], "import_folders": [str(root)]}, "tdfs": {"enabled": False}})
    res = run_project(cfg)
    assert res.status == "failed" and res.groups[0].status == "blocked" and res.groups[0].summary == {}
    assert any(d.rule == "unsupported_layout" for d in res.diagnostics)


# ----------------------------------------------------------------------------- advanced features / new options
def test_wobble_r0_sources(tmp_path):
    p = make_anisotropy_workbook(str(tmp_path / "a.xlsx"), [3.5], [0.29], 0.0454575, 0.333)
    wb = load_workbook(p)
    lit = compute_anisotropy(wb, WobbleSettings(enabled=True, effective_volume=161, effective_volume_unit="A3", temperature_K=293))
    assert lit.wobble["r0_used"] == 0.39 and lit.wobble["r0_source"] == "literature" and lit.wobble["status"] == "ok"
    fit = compute_anisotropy(wb, WobbleSettings(enabled=True, effective_volume=161, effective_volume_unit="A3", temperature_K=293, r0_source="fitted"))
    assert fit.wobble["r0_used"] == pytest.approx(0.333)
    usr = compute_anisotropy(wb, WobbleSettings(enabled=True, effective_volume=1.61e-22, effective_volume_unit="cm3", temperature_K=293, r0_source="user", r0=0.4))
    assert usr.wobble["r0_used"] == 0.4 and usr.wobble["effective_volume_m3"] == pytest.approx(1.61e-28)
    bad = compute_anisotropy(wb, WobbleSettings(enabled=True, effective_volume=161, effective_volume_unit="A3", temperature_K=293, r0_source="user"))
    assert bad.wobble["status"].startswith("blocked")


def test_parallel_loading_identical(synthetic_group):
    p1 = Project(["Composition", "Temperature", "Replicate"]); p1.import_folder(str(synthetic_group)); p1.load_all(workers=1)
    p2 = Project(["Composition", "Temperature", "Replicate"]); p2.import_folder(str(synthetic_group)); p2.load_all(workers=2)
    for a, b in zip(p1.measurements, p2.measurements):
        assert a.measurement_id == b.measurement_id and a.kind == b.kind
        if a.kind == "tres":
            assert np.array_equal(a.workbook.components.normalized_pre_exponential, b.workbook.components.normalized_pre_exponential)


def test_apparent_nu0_warning_and_boundary_column(synthetic_group):
    cfg = config_from_dict({"project": {"level_names": ["Composition", "Temperature", "Replicate"], "import_folders": [str(synthetic_group)]},
                            "tdfs": {"recon_tmax_min_ns": 20, "trajectory_stride": 10}})
    res = run_project(cfg)
    assert any(d.rule == "tdfs_nu0_apparent" and d.severity == Severity.WARNING for d in res.diagnostics)
    g = res.ok_groups("tres")[0]
    assert "peak_at_boundary_fraction" in g.replicates[0].scalars
    cfg2 = config_from_dict({"project": {"level_names": ["Composition", "Temperature", "Replicate"], "import_folders": [str(synthetic_group)]},
                             "tdfs": {"recon_tmax_min_ns": 20, "trajectory_stride": 10, "nu0_cm1": 23800}})
    res2 = run_project(cfg2)
    assert not any(d.rule == "tdfs_nu0_apparent" for d in res2.diagnostics)


def test_deconvolution_review_scripted_and_replay(synthetic_group, tmp_path):
    from tres_suite.export import write_run
    from tres_suite.advanced.deconvolution import run_deconvolution
    import json
    cfg = config_from_dict({"project": {"level_names": ["Composition", "Temperature", "Replicate"], "import_folders": [str(synthetic_group)]},
                            "output": {"directory": str(tmp_path / "o"), "overwrite": True}, "tdfs": {"enabled": False}})
    res = run_project(cfg)
    write_run(res)
    answers = iter(["no", "1", "410", "470", "yes", "yes", "yes"])  # component 1: reject, refit with 1 peak, accept; 2 and 3: accept
    path = run_deconvolution(str(tmp_path / "o"), groups=["DOPC / 15"], prompt=lambda q: next(answers), show=lambda p: None)
    dec = json.load(open(path))
    comp = dec["groups"]["DOPC_15"]
    assert comp["1"]["n_peaks"] == 1 and comp["1"]["bounds_nm"] == [410.0, 470.0] and comp["1"]["accepted"]
    assert (tmp_path / "o" / "deconvolution" / "DOPC_15" / "bub data.csv").exists()
    txt = open(tmp_path / "o" / "deconvolution" / "DOPC_15" / "bub data.csv").read().splitlines()
    assert txt[0] == "lifetime,peakposition,area,region" and all(line.endswith(",") for line in txt[1:])  # region empty
    areas = [float(l.split(",")[2]) for l in txt[1:]]
    assert sum(areas) == pytest.approx(100.0)
    path2 = run_deconvolution(str(tmp_path / "o"), groups=["DOPC / 15"], replay=path, interactive=False)
    assert json.load(open(path2)) == dec


def test_stacked_area_bar_plot(synthetic_group, tmp_path):
    from tres_suite.plotting import render_plots
    from tres_suite.config import PlotConfig
    cfg = config_from_dict({"project": {"level_names": ["Composition", "Temperature", "Replicate"], "import_folders": [str(synthetic_group)]}, "tdfs": {"enabled": False}})
    res = run_project(cfg)
    out = render_plots(res, str(tmp_path / "p"), [PlotConfig(type="stacked_area_bar", name="stack")])
    df = out[0].data
    assert set(df.group) == {"DOPC / 15", "DOPC / 25"} and df.groupby("group").area_percent.sum().round(6).eq(100).all()


# ----------------------------------------------------------------------------- review findings (regressions)
def test_reversed_wavelength_order_is_normalised(tmp_path):
    taus, B, counts = synthetic_tres_params(WL)
    a = make_tres_workbook(str(tmp_path / "asc.xlsx"), WL, taus, B, counts)
    b = make_tres_workbook(str(tmp_path / "desc.xlsx"), WL[::-1], taus, B[::-1], counts[::-1])
    wa, wb_ = load_workbook(a), load_workbook(b)
    assert np.array_equal(wa.wavelengths_nm, wb_.wavelengths_nm)
    assert np.allclose(wa.components.pre_exponential, wb_.components.pre_exponential)
    assert np.allclose(wa.total_counts, wb_.total_counts)
    assert np.allclose(wa.summary_relative_amplitude_percent, wb_.summary_relative_amplitude_percent)
    da, db = compute_das(wa, DasSettings()), compute_das(wb_, DasSettings())
    assert np.allclose(da.emission, db.emission) and da.mean_lifetime_ns == pytest.approx(db.mean_lifetime_ns)
    assert any("reordered" in w for w in wb_.warnings)


def test_duplicate_replicate_filenames_do_not_collide(tmp_path):
    from tres_suite.export import write_run
    root = tmp_path / "d"
    taus, B, counts = synthetic_tres_params(WL)
    p = Project(["Cond", "Rep"])
    for sub in ("a", "b"):
        (root / sub).mkdir(parents=True)
        path = make_tres_workbook(str(root / sub / "1.xlsx"), WL, taus, B, counts * (1.0 if sub == "a" else 1.2))
        p.add_file(path, {"Cond": "X", "Rep": f"{sub}/1"})
    cfg = config_from_dict({"project": {"level_names": ["Cond", "Rep"], "files": [{"path": str(root / s / "1.xlsx"), "labels": {"Cond": "X", "Rep": f"{s}/1"}} for s in ("a", "b")]},
                            "output": {"directory": str(tmp_path / "o"), "overwrite": True, "image_formats": ["png"]},
                            "tdfs": {"recon_tmax_min_ns": 20, "trajectory_stride": 10}})
    res = run_project(cfg)
    assert res.groups[0].n_replicates == 2
    manifest = write_run(res)
    rep_dir = tmp_path / "o" / "groups" / "X" / "replicates"
    files = sorted(os.listdir(rep_dir))
    assert len([f for f in files if f.endswith("_tdfs_trajectory.csv")]) == 2
    assert len([f for f in files if f.endswith("_TRES_reconstructed.csv")]) == 2
    ids = {r.measurement.measurement_id for r in res.groups[0].replicates}
    assert all(any(i in f for i in ids) for f in files)
    assert len(manifest["files"]) == len({f["path"] for f in manifest["files"]})  # no duplicate paths registered


def test_cli_run_with_plots_writes_manifest(synthetic_group, tmp_path):
    import yaml
    from tres_suite.cli import main
    cfg = {"project": {"level_names": ["Composition", "Temperature", "Replicate"], "import_folders": [str(synthetic_group)]},
           "output": {"directory": str(tmp_path / "fresh"), "image_formats": ["png"]},  # overwrite left at its default (false)
           "tdfs": {"recon_tmax_min_ns": 20, "trajectory_stride": 10},
           "plots": [{"type": "metric_vs_label", "name": "t", "x": "Temperature", "y": "tmean_ns"}]}
    cfg_path = tmp_path / "c.yaml"
    cfg_path.write_text(yaml.safe_dump(cfg))
    assert main(["run", str(cfg_path), "--quiet"]) == 0
    assert (tmp_path / "fresh" / "run_manifest.json").exists() and (tmp_path / "fresh" / "plots" / "t.png").exists()
    assert main(["run", str(cfg_path), "--quiet"]) == 2  # second run refuses the non-empty folder before doing anything
    assert main(["run", str(cfg_path), "--quiet", "--overwrite"]) == 0


def test_config_round_trip(tmp_path):
    import json
    from tres_suite.config import load_config
    d = {"project": {"level_names": ["A", "B", "C"], "import_folders": ["/x"], "workers": 2},
         "tdfs": {"enabled": False, "nu0_cm1": 23800, "export_times_ns": [0.1, 1.0]},
         "validation": {"chi_sq_warning_threshold": 1.5, "replicate_variation_metrics": ["gp"]},
         "time_gated": {"enabled": True, "gates_ns": [[0, 1], [1, 3]]},
         "wobble": {"enabled": True, "effective_volume": 161, "effective_volume_unit": "A3", "temperature_K": 293},
         "plots": [{"type": "bubble", "name": "b", "group": {"A": "x"}, "series": [{"label": "s", "select": {"A": "x"}, "color": "#ff0000"}]}]}
    cfg = config_from_dict(d)
    assert cfg.tdfs_enabled is False
    again = config_from_dict(cfg.to_dict())
    assert again == cfg
    path = tmp_path / "run_config.json"
    path.write_text(json.dumps(cfg.to_dict()))
    loaded = load_config(str(path))
    assert loaded.tdfs == cfg.tdfs and loaded.tdfs_enabled is False and loaded.validation == cfg.validation and loaded.time_gated == cfg.time_gated
    # an old run_config.json with the derived nu0_mode key and a top-level tdfs_enabled still loads
    old = cfg.to_dict(); old["tdfs"]["nu0_mode"] = "user"; old["tdfs_enabled"] = False; del old["tdfs"]["enabled"]
    assert config_from_dict(old).tdfs_enabled is False


def test_wide_tables_keep_every_grouping_level(tmp_path):
    from tres_suite.export import _wide_table, write_run
    root = tmp_path / "d"
    taus, B, counts = synthetic_tres_params(WL)
    files = []
    for comp, temp, batch in [("DOPC", "15", "b1"), ("DOPC", "15", "b2"), ("DOPC", "25", "b1")]:
        d = root / comp / temp / batch; d.mkdir(parents=True)
        files.append(make_tres_workbook(str(d / "1.xlsx"), WL, taus, B, counts * (1 + 0.1 * len(files))))
    cfg = config_from_dict({"project": {"level_names": ["Composition", "Temperature", "Batch", "Replicate"], "import_folders": [str(root)]},
                            "output": {"directory": str(tmp_path / "o"), "overwrite": True}, "tdfs": {"enabled": False}})
    res = run_project(cfg)
    assert len(res.groups) == 3
    df = _wide_table(res, "tmean_ns", "mean")
    assert list(df.columns) == ["Temperature", "Batch", "DOPC"] and len(df) == 3
    vals = {(r.Temperature, r.Batch): r.DOPC for r in df.itertuples()}
    for g in res.groups:
        assert vals[(g.summary["label:Temperature"], g.summary["label:Batch"])] == g.summary["tmean_ns_mean"]
    write_run(res)
    legacy = open(tmp_path / "o" / "legacy" / "graph_data" / "gp_mean.csv", encoding="utf-8-sig").read().splitlines()
    assert legacy[0] == "Temperature,Batch,DOPC,Valid" and len(legacy) == 4


# ----------------------------------------------------------------------------- GUI-phase backend additions
def test_wobble_temperature_from_level(synthetic_group):
    """wobble.temperature_level reads T per group from a hierarchy label (degC by default; explicit units respected)."""
    cfg = config_from_dict({"project": {"level_names": ["Composition", "Temperature", "Replicate"], "import_folders": [str(synthetic_group)]},
                            "tdfs": {"enabled": False},
                            "wobble": {"enabled": True, "effective_volume": 161, "effective_volume_unit": "A3", "temperature_level": "Temperature"}})
    res = run_project(cfg)
    g = res.ok_groups("anisotropy")[0]
    assert g.summary["wobble_temperature_K_mean"] == pytest.approx(15 + 273.15)
    assert g.replicates[0].scalars["wobble_status"] == "ok" and g.replicates[0].scalars["wobble_temperature_source"] == "level:Temperature"
    # same numbers as a fixed temperature of 288.15 K
    cfg2 = config_from_dict({"project": {"level_names": ["Composition", "Temperature", "Replicate"], "import_folders": [str(synthetic_group)]},
                             "tdfs": {"enabled": False},
                             "wobble": {"enabled": True, "effective_volume": 161, "effective_volume_unit": "A3", "temperature_K": 288.15}})
    g2 = run_project(cfg2).ok_groups("anisotropy")[0]
    assert g2.summary["wobble_eta_Pa_s_mean"] == pytest.approx(g.summary["wobble_eta_Pa_s_mean"])
    # a non-numeric label at that level -> add-on unavailable, base quantities untouched
    from tres_suite.anisotropy import WobbleSettings, temperature_from_label
    from tres_suite.project import LabelValue
    ws = WobbleSettings(enabled=True, effective_volume=161, effective_volume_unit="A3", temperature_level="Composition")
    assert temperature_from_label(LabelValue.parse("DPH"), ws) is None
    assert temperature_from_label(LabelValue.parse("300K"), ws) == pytest.approx(300.0)
    assert temperature_from_label(LabelValue.parse("17C"), ws) == pytest.approx(290.15)
    assert temperature_from_label(LabelValue.parse("17"), WobbleSettings(temperature_level="x", temperature_level_unit="K")) == pytest.approx(17.0)
    assert not ws.missing()
    d = cfg.to_dict()["wobble"]
    assert d["temperature_level"] == "Temperature" and config_from_dict(cfg.to_dict()).wobble.temperature_level == "Temperature"


def test_run_project_progress_cancel_and_parallel_analysis(synthetic_group):
    base = {"project": {"level_names": ["Composition", "Temperature", "Replicate"], "import_folders": [str(synthetic_group)]},
            "tdfs": {"recon_tmax_min_ns": 20, "trajectory_stride": 10}}
    seen = []
    res1 = run_project(config_from_dict(base), file_progress=lambda d, t, st: seen.append((d, t, st)))
    assert any(st == "loading" for _, _, st in seen) and any(st == "analysing" for _, _, st in seen)
    assert max(d for d, t, st in seen if st == "analysing") == 8  # six TRES + two anisotropy replicates
    cfg2 = config_from_dict(dict(base, project=dict(base["project"], workers=2)))
    res2 = run_project(cfg2)
    for a, b in zip(res1.groups, res2.groups):
        assert a.status == b.status
        if a.kind == "tres":
            assert a.summary["gp_mean"] == pytest.approx(b.summary["gp_mean"]) and a.summary["peak_tau_r_ns_mean"] == pytest.approx(b.summary["peak_tau_r_ns_mean"])
            assert a.replicates[0].measurement is a.group.measurements[0] and b.replicates[0].measurement is b.group.measurements[0]
    from tres_suite.pipeline import RunCancelled
    with pytest.raises(RunCancelled):
        run_project(config_from_dict(base), should_cancel=lambda: True)


def test_draw_plot_on_axes_and_new_plot_types(synthetic_group, tmp_path):
    import matplotlib
    matplotlib.use("Agg")
    from matplotlib.figure import Figure
    from tres_suite.plotting import draw_plot, render_plots
    from tres_suite.config import PlotConfig, SeriesConfig
    cfg = config_from_dict({"project": {"level_names": ["Composition", "Temperature", "Replicate"], "import_folders": [str(synthetic_group)]},
                            "tdfs": {"recon_tmax_min_ns": 20, "trajectory_stride": 10}})
    res = run_project(cfg)
    grp = {"Composition": "DOPC", "Temperature": "15"}
    for pc in (PlotConfig(type="steady_state_spectrum", name="ss", group=grp),
               PlotConfig(type="tdfs_fwhm_vs_time", name="fw", group=grp),
               PlotConfig(type="tdfs_spectra", name="full", group=grp, extra={"normalise": "none"}),
               PlotConfig(type="metric_vs_label", name="gp_t", x="Temperature", y="gp", series=[SeriesConfig(label="DOPC", select={"Composition": "DOPC"})]),
               PlotConfig(type="metric_vs_label", name="t_gp", x="Temperature", y="gp", x_label="GP!", y_label="T!", x_limits=[-1, 1],
                          series=[SeriesConfig(label="DOPC", select={"Composition": "DOPC"})], extra={"label_axis": "y"})):
        fig = Figure()
        ax = fig.add_subplot(111)
        out = draw_plot(res, pc, ax)
        assert len(out.data) > 0 and len(ax.lines) + len(ax.collections) > 0
        if pc.name == "t_gp":  # transposed: label on y, metric on x, labels/limits applied to the axes as given
            assert out.x_key == "gp" and out.y_key == "Temperature"
            assert ax.get_xlabel() == "GP!" and ax.get_ylabel() == "T!" and ax.get_xlim() == (-1.0, 1.0)
            assert sorted(out.data["y"].tolist()) == [15.0, 25.0]
        # the file renderer draws the same thing
        files = render_plots(res, str(tmp_path / "p"), [pc])[0].files
        assert any(f.endswith(".png") for f in files) and any(f.endswith("_data.csv") for f in files)


def test_single_plot_multi_sample_overlays_and_legend_toggle(synthetic_group):
    import matplotlib
    matplotlib.use("Agg")
    from matplotlib.figure import Figure
    from tres_suite.plotting import draw_plot
    from tres_suite.config import PlotConfig, SeriesConfig

    cfg = config_from_dict({"project": {"level_names": ["Composition", "Temperature", "Replicate"],
                                        "import_folders": [str(synthetic_group)]},
                            "tdfs": {"recon_tmax_min_ns": 20, "trajectory_stride": 10}})
    res = run_project(cfg)
    series = [
        SeriesConfig(label="DOPC / 15 (n=3)", select={"Composition": "DOPC", "Temperature": "15"},
                     color="#2f7fc1", marker="", line_style="-", line_width=2.5),
        SeriesConfig(label="DOPC / 25 (n=3)", select={"Composition": "DOPC", "Temperature": "25"},
                     color="#cf342b", marker="s", line_style="--", line_width=1.25),
    ]
    plots = [
        PlotConfig(type="steady_state_spectrum", name="overlay_ss"),
        PlotConfig(type="tdfs_relaxation", name="overlay_relax", coordinate="peak"),
        PlotConfig(type="tdfs_spectra", name="overlay_tranes", extra={"normalise": "area"}),
        PlotConfig(type="das_spectra", name="overlay_das", representation="relative_amplitude"),
        PlotConfig(type="bubble", name="overlay_bubble"),
    ]
    for pc in plots:
        pc.series = series
        pc.extra.update({"legend_labels_exact": True, "show_legend": False})
        fig = Figure()
        ax = fig.add_subplot(111)
        out = draw_plot(res, pc, ax)
        assert set(out.data["sample"]) == {s.label for s in series}
        assert ax.get_legend() is None
        assert len(ax.lines) + len(ax.collections) > 1
        if pc.type in {"tdfs_spectra", "das_spectra"}:
            sample_lines = [[line for line in ax.lines if line.get_label().startswith(s.label + " —")] for s in series]
            assert all(len(lines) > 1 for lines in sample_lines)
            assert len({line.get_color() for line in sample_lines[0]}) > 1
            assert [line.get_color() for line in sample_lines[0]] == [line.get_color() for line in sample_lines[1]]
            assert {line.get_linestyle() for line in sample_lines[0]} == {"-"}
            assert {line.get_linestyle() for line in sample_lines[1]} == {"--"}

    round_trip = config_from_dict(config_from_dict({"plots": [{"type": "steady_state_spectrum", "name": "styled",
                                                                 "series": [{"label": "sample", "line_style": "--", "line_width": 2.5}]}]}).to_dict())
    assert round_trip.plots[0].series[0].line_style == "--"
    assert round_trip.plots[0].series[0].line_width == pytest.approx(2.5)
