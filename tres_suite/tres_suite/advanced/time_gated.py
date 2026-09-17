"""Time-gated spectra and IRF-limited "measured" energy metrics (advanced feature).

Port of sections 1, 4, 4a and 4b of ``single laurdan.py``.  These use the raw
decays of the ``Data`` sheet (``Time / s``, ``<wl>-DECAY`` columns), are limited
by the ~250 ps instrument response and are **not** used for the TDFS
parameters (those come from the reconstruction).  They are provided for
comparison only and must be enabled explicitly.

Per replicate: decays are normalised by their total photon sum, time-shifted so
that the summed-intensity maximum sits at t = 0 and linearly interpolated onto
a common grid (default −2 … 50 ns, 0.05 ns).  Replicates are averaged
point-wise, then smoothed with a Gaussian (σ = 5 time steps, 1 wavelength step,
legacy values).  From the smoothed surface: area-normalised time-gated spectra
for the configured gates, and for t ≥ t_peak while the summed intensity is
above ``floor_fraction`` of its maximum the wavenumber-domain peak, centre of
mass and FWHM (same spline method as the reconstruction), the peak
trajectory v(t) truncated at its first minimum and C(t) = (v − v_∞)/(v_0 − v_∞)
with v_0/v_∞ the means of the first/last three points (legacy convention).
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field, asdict
from typing import Any, Dict, List, Optional, Sequence, Tuple

import numpy as np
import openpyxl
from scipy.ndimage import gaussian_filter1d

from ..tdfs import _make_grid, spectral_moments

_trapz = getattr(np, "trapezoid", None) or np.trapz  # type: ignore[attr-defined]


@dataclass
class TimeGatedSettings:
    enabled: bool = False
    common_time_ns: Tuple[float, float, float] = (-2.0, 50.0, 0.05)  # start, stop, step
    smooth_sigma_time_steps: float = 5.0
    smooth_sigma_wavelength_steps: float = 1.0
    gates_ns: Sequence[Sequence[float]] = ((0.0, 0.5), (0.5, 2.0), (2.0, 7.0))
    gate_mode: str = "integral"  # integral | mean
    floor_fraction: float = 0.05
    fine_nu_points: int = 1000

    def to_dict(self) -> Dict[str, Any]:
        d = asdict(self)
        d["gates_ns"] = [list(g) for g in self.gates_ns]
        return d


def read_data_sheet(path: str) -> Tuple[np.ndarray, np.ndarray, np.ndarray]:
    """Return (time_ns, wavelengths_nm, decays (n_t, n_wl)) from the ``Data`` sheet."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        if "Data" not in wb.sheetnames:
            raise ValueError("no 'Data' sheet")
        rows = wb["Data"].iter_rows(values_only=True)
        header = [("" if c is None else str(c).strip()) for c in next(rows)]
        t_col = next((i for i, h in enumerate(header) if h.lower().startswith("time")), None)
        if t_col is None:
            raise ValueError("no time column in 'Data' sheet")
        cols = []
        for i, h in enumerate(header):
            m = re.match(r"^\s*([-+]?\d+(?:\.\d+)?)\s*-\s*DECAY\s*$", h, re.IGNORECASE)
            if m:
                cols.append((float(m.group(1)), i))
        if not cols:
            raise ValueError("no '<wl>-DECAY' columns in 'Data' sheet")
        cols.sort()
        t, data = [], []
        for r in rows:
            if r is None or t_col >= len(r) or r[t_col] is None:
                break
            try:
                tv = float(r[t_col])
            except (TypeError, ValueError):
                break
            t.append(tv)
            data.append([float(r[i]) if i < len(r) and r[i] is not None else np.nan for _, i in cols])
    finally:
        wb.close()
    t = np.asarray(t, float)
    unit = 1e9 if np.nanmax(np.abs(t)) < 1e-3 else 1.0  # seconds -> ns
    return t * unit, np.array([w for w, _ in cols], float), np.asarray(data, float)


def _align_and_interpolate(t_ns: np.ndarray, data: np.ndarray, common: np.ndarray) -> np.ndarray:
    total = np.nansum(data, axis=1)
    t0 = t_ns[int(np.nanargmax(total))]
    denom = np.nansum(data)
    if not denom > 0:
        raise ValueError("total intensity is not positive")
    norm = data / denom
    out = np.zeros((len(common), data.shape[1]))
    for w in range(data.shape[1]):
        out[:, w] = np.interp(common, t_ns - t0, norm[:, w], left=0.0, right=0.0)
    return out


def analyse_time_gated(paths: Sequence[str], settings: TimeGatedSettings) -> Dict[str, Any]:
    """Average the replicate decay surfaces and compute the legacy measured metrics."""
    a, b, step = settings.common_time_ns
    common = np.arange(a, b, step)
    wl_ref = None
    mats = []
    for p in paths:
        t_ns, wl, data = read_data_sheet(p)
        if wl_ref is None:
            wl_ref = wl
        elif len(wl) != len(wl_ref) or not np.allclose(wl, wl_ref):
            raise ValueError(f"{p}: Data-sheet wavelengths differ from the first replicate")
        mats.append(_align_and_interpolate(t_ns, data, common))
    avg = np.mean(np.array(mats), axis=0)
    smooth = gaussian_filter1d(avg, sigma=settings.smooth_sigma_time_steps, axis=0)
    smooth = gaussian_filter1d(smooth, sigma=settings.smooth_sigma_wavelength_steps, axis=1)
    wl = wl_ref
    out: Dict[str, Any] = {"time_ns": common, "wavelengths_nm": wl, "avg_matrix": avg, "avg_matrix_smooth": smooth, "n_replicates": len(mats),
                           "settings": settings.to_dict()}
    total = smooth.sum(axis=1)
    ipk = int(np.argmax(total))
    thresh = settings.floor_fraction * total[ipk]
    # gates
    gates, labels = [], []
    order = np.argsort(wl)
    for lo, hi in settings.gates_ns:
        mask = (common >= lo) & (common <= hi)
        if not mask.any():
            spec = np.zeros(len(wl))
        elif settings.gate_mode == "mean":
            spec = smooth[mask].mean(axis=0)
        else:
            spec = _trapz(smooth[mask], common[mask], axis=0)
        area = float(_trapz(spec[order], wl[order]))
        gates.append(spec / area if area > 0 else spec)
        labels.append(f"{lo:g}-{hi:g} ns")
    out["gated_spectra"] = np.array(gates)
    out["gate_labels"] = labels
    # measured energy metrics
    grid = _make_grid(wl, settings.fine_nu_points)
    tt, pk, cm, fw = [], [], [], []
    for i in range(ipk, len(common)):
        if total[i] < thresh:
            break
        p_, c_, f_, _, _ = spectral_moments(smooth[i], grid, True)
        tt.append(common[i]); pk.append(p_); cm.append(c_); fw.append(f_)
    tt, pk, cm, fw = map(lambda x: np.asarray(x, float), (tt, pk, cm, fw))
    out.update({"metric_time_ns": tt, "peak_cm-1": pk, "com_cm-1": cm, "fwhm_cm-1": fw})
    # v(t) / C(t) with first-minimum truncation (legacy)
    nonneg = tt >= 0
    tv, vv = tt[nonneg], pk[nonneg]
    if len(vv) >= 3:
        imin = int(np.nanargmin(vv))
        tv, vv = tv[:imin + 1], vv[:imin + 1]
        if len(vv) >= 3:
            v0 = float(np.mean(vv[:3])); vinf = float(np.mean(vv[-3:]))
            den = v0 - vinf
            c = (vv - vinf) / den if abs(den) > 1e-9 else np.full_like(vv, np.nan)
        else:
            c = np.full_like(vv, np.nan)
    else:
        c = np.full_like(vv, np.nan)
    out.update({"ct_time_ns": tv, "v_t_cm-1": vv, "C_t": c})
    return out
